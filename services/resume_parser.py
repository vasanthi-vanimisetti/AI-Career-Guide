"""Extract readable text from supported resume formats."""

from pathlib import Path

import pdfplumber
from docx import Document


class ResumeParseError(ValueError):
    """Raised when a resume cannot be extracted safely."""


def extract_resume_text(path: Path) -> str:
    """Extract text from a PDF or DOCX document."""
    try:
        if path.suffix.lower() == ".pdf":
            with pdfplumber.open(path) as pdf:
                return "\n".join(page.extract_text() or "" for page in pdf.pages)
        if path.suffix.lower() == ".docx":
            document = Document(path)
            paragraphs = [paragraph.text for paragraph in document.paragraphs]
            paragraphs.extend(cell.text for table in document.tables for row in table.rows for cell in row.cells)
            return "\n".join(paragraphs)
        raise ResumeParseError("Unsupported file type. Upload a PDF or DOCX file.")
    except ResumeParseError:
        raise
    except Exception as error:
        raise ResumeParseError("We could not read this resume. Check that the file is not damaged or password-protected.") from error